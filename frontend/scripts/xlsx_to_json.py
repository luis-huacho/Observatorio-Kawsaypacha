#!/usr/bin/env python3
"""Convierte data/Base_Nivel Peligro_CCPP_Cusco.xlsx a JSON normalizados.

Genera:
  prototype/public/data/ccpp.json      — centros poblados deduplicados
  prototype/public/data/peligros.json  — clasificaciones (long format)

Correr desde la raíz del repo:
  python3 prototype/scripts/xlsx_to_json.py
"""
from pathlib import Path
import json
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "data" / "Base_Nivel Peligro_CCPP_Cusco.xlsx"
OUT = ROOT / "prototype" / "public" / "data"
OUT.mkdir(parents=True, exist_ok=True)


def ubigeo_distrito(codigo_ccpp: str) -> str:
    return codigo_ccpp[:6] if codigo_ccpp and len(codigo_ccpp) >= 6 else ""


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ccpp_by_code: dict[str, dict] = {}
    peligros: list[dict] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = ws.iter_rows(min_row=2, values_only=True)
        for r in rows:
            if not r or not r[3]:
                continue
            (
                depto, prov, dist, codigo, nombre, categoria,
                altitud, lon, lat, poblacion,
                peligro, tip_pelig, nivel_peli, fuente, link,
            ) = r[:15]

            codigo = str(codigo).strip()
            if codigo not in ccpp_by_code:
                ccpp_by_code[codigo] = {
                    "codigo": codigo,
                    "nombre": (nombre or "").strip(),
                    "categoria": (categoria or "").strip(),
                    "departamento": (depto or "").strip(),
                    "provincia": (prov or "").strip(),
                    "distrito": (dist or "").strip(),
                    "ubigeo_distrito": ubigeo_distrito(codigo),
                    "lat": float(lat) if lat is not None else None,
                    "lon": float(lon) if lon is not None else None,
                    "altitud": int(altitud) if altitud is not None else None,
                    "poblacion": int(poblacion) if poblacion is not None else None,
                }

            if nivel_peli is not None and peligro:
                peligros.append({
                    "codigo_ccpp": codigo,
                    "peligro": str(peligro).strip(),
                    "tipo": (str(tip_pelig).strip() if tip_pelig else None),
                    "nivel": int(nivel_peli),
                    "fuente": (str(fuente).strip() if fuente else None),
                    "fuente_url": (str(link).strip() if link else None),
                })

    ccpp = sorted(ccpp_by_code.values(), key=lambda x: x["codigo"])

    (OUT / "ccpp.json").write_text(json.dumps(ccpp, ensure_ascii=False), encoding="utf-8")
    (OUT / "peligros.json").write_text(json.dumps(peligros, ensure_ascii=False), encoding="utf-8")

    print(f"ccpp.json: {len(ccpp)} centros poblados")
    print(f"peligros.json: {len(peligros)} clasificaciones de peligro")


if __name__ == "__main__":
    main()
